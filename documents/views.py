import csv
import json
import logging
import re

import requests

from django.conf import settings
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone

logger = logging.getLogger(__name__)

_PATRON_CLE_URL = re.compile(r'[?&](?:key|api[_-]?key)=[^&\s]+')

from clients.models import Client
from dossiers.models import Dossier
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from .forms import DocumentForm, GenerationDocumentForm
from .gabarits import GABARIT_MISE_EN_DEMEURE
from .ia import generer_mise_en_demure, analyser_document as analyser_document_ia, extraire_texte
from .models import Document, DocumentGenere, JournalDocumentGenere, AnalyseDocument
from facturation.services import quota_restant, enregistrer_usage


def _message_erreur_securise(exc):
    """Message d'erreur affichable à l'utilisateur, sans jamais exposer de clé API.

    Le détail complet (avec l'URL et la clé) est journalisé côté serveur.
    """
    logger.error("Erreur applicative : %s", exc)
    texte = str(exc or '')
    if not texte:
        return "Une erreur est survenue. Réessayez."
    return _PATRON_CLE_URL.sub('[clé masquée]', texte)


def _message_ia_securise(exc):
    """Message utilisateur pour les appels IA, avec traitement des timeouts réseau."""
    if isinstance(exc, (requests.exceptions.Timeout, requests.exceptions.ConnectionError)):
        logger.error("Erreur IA (réseau) : %s", exc)
        return "Le service IA n'a pas répondu à temps (réseau). Réessayez dans quelques instants."
    return _message_erreur_securise(exc)

@login_required
def liste_documents(request):
    qs = (Document.objects.select_related('dossier', 'client', 'auteur')
          .filter(Q(dossier__avocat=request.user) | Q(auteur=request.user)))

    q = request.GET.get('q', '').strip()
    dossier_id = request.GET.get('dossier', '').strip()
    type_f = request.GET.get('type', '').strip()
    categorie = request.GET.get('categorie', '').strip()
    tri = request.GET.get('tri', 'date')

    if q:
        qs = qs.filter(nom_original__icontains=q)
    if dossier_id:
        qs = qs.filter(dossier_id=dossier_id)
    if type_f:
        qs = qs.filter(type_fichier=type_f)
    if categorie:
        qs = qs.filter(categorie=categorie)

    if tri == 'nom':
        qs = qs.order_by('nom_original')
    elif tri == 'taille':
        qs = qs.order_by('-taille_octets')
    else:
        qs = qs.order_by('-date_upload')

    if request.GET.get('export'):
        return exporter_csv(request, qs)

    paginator = Paginator(qs, 10)
    documents = paginator.get_page(request.GET.get('page', 1))

    form = DocumentForm(request.POST or None, request.FILES or None, user=request.user)
    if request.method == 'POST' and form.is_valid():
        doc = form.save(commit=False)
        doc.auteur = request.user
        doc.nom_original = doc.fichier.name
        doc.taille_octets = doc.fichier.size
        doc.type_fichier = doc.fichier.name.lower().rsplit('.', 1)[-1]
        doc.save()
        return redirect('liste_documents')

    return render(request, 'documents_liste.html', {
        'documents': documents,
        'form': form,
        'q': q,
        'dossier_id': dossier_id,
        'type_f': type_f,
        'categorie': categorie,
        'tri': tri,
        'dossiers': Dossier.objects.filter(avocat=request.user).order_by('titre'),
        'clients': Client.objects.filter(dossiers__avocat=request.user).distinct().order_by('nom'),
        'types': settings.UPLOAD_EXTENSIONS,
        'categories': Document.CATEGORIE_CHOICES,
        'active': 'documents',
    })


def exporter_csv(request, qs):
    response = HttpResponse(content_type='text/csv')
    response['Content-Disposition'] = 'attachment; filename="documents.csv"'
    w = csv.writer(response)
    w.writerow(['Nom', 'Dossier', 'Contact', 'Catégorie', 'Date', 'Taille (octets)', 'Type', 'Auteur'])
    for d in qs:
        w.writerow([d.nom_original, d.dossier.titre if d.dossier else '',
                    d.client.nom if d.client else '', d.get_categorie_display(),
                    d.date_upload.strftime('%Y-%m-%d %H:%M'),
                    d.taille_octets or '', d.type_fichier,
                    (d.auteur.get_full_name() or d.auteur.username) if d.auteur else ''])
    return response

def _valeurs_gabarit(request, dossier, form):
    client = dossier.client
    profil = getattr(request.user, 'profil', None)
    avocat = request.user.get_full_name() or request.user.username
    return {
        'cabinet': 'Cabinet',
        'telephone_cabinet': profil.telephone if profil and profil.telephone else '',
        'email_cabinet': request.user.email or '',
        'date': timezone.localdate().strftime('%d %B %Y'),
        'nom_destinataire': form.cleaned_data.get('nom_destinataire', ''),
        'adresse_destinataire': form.cleaned_data.get('adresse_destinataire', ''),
        'objet': dossier.titre,
        'nom_client': str(client),
        'adresse_client': form.cleaned_data.get('adresse_client', ''),
        'faits_resume': dossier.notes or '',
        'montant_reclame': str(form.cleaned_data.get('montant_reclame', '')),
        'delai_jours': form.cleaned_data.get('delai_jours', 10),
        'avocat': avocat,
        'signature_cabinet': 'Cabinet',
    }


@login_required
def generer_document(request):
    form = GenerationDocumentForm(request.POST or None, user=request.user)
    if request.method == 'POST' and form.is_valid():
        dossier = form.cleaned_data['dossier']
        valeurs = _valeurs_gabarit(request, dossier, form)
        gabarit_rempli = GABARIT_MISE_EN_DEMEURE.format(**valeurs)
        if quota_restant(request.user) <= 0:
            messages.error(request, "Quota IA mensuel dépassé. Contactez l'administrateur.")
            return render(request, 'documents_generation.html', {'form': form, 'active': 'documents_generes'})
        try:
            contenu = generer_mise_en_demure(gabarit_rempli, valeurs)
        except Exception as e:
            messages.error(request, "Erreur lors de la génération IA : {}".format(_message_ia_securise(e)))
            return render(request, 'documents_generation.html', {'form': form, 'active': 'documents_generes'})
        enregistrer_usage(request.user, 'generation', octets=len(contenu))
        doc = DocumentGenere.objects.create(
            dossier=dossier,
            type_document='mise_en_demure',
            contenu_brouillon=contenu,
            statut='brouillon',
            parametres_generation=json.dumps({
                'gabarit_rempli': gabarit_rempli,
                'valeurs': valeurs,
            }, ensure_ascii=False),
        )
        JournalDocumentGenere.objects.create(document=doc, utilisateur=request.user, action='generation')
        messages.success(request, "Brouillon généré. Révisez-le avant validation.")
        return redirect('reviser_document', pk=doc.pk)
    return render(request, 'documents_generation.html', {'form': form, 'active': 'documents_generes'})


@login_required
def reviser_document(request, pk):
    doc = get_object_or_404(
        DocumentGenere.objects.select_related('dossier', 'dossier__client', 'valide_par'),
        pk=pk, dossier__avocat=request.user)
    if request.method == 'POST':
        action = request.POST.get('action')
        if action == 'enregistrer':
            doc.contenu_brouillon = request.POST.get('contenu', '').strip()
            doc.save()
            JournalDocumentGenere.objects.create(document=doc, utilisateur=request.user, action='modification')
            messages.success(request, "Brouillon enregistré.")
        elif action == 'valider':
            if doc.statut != 'envoye':
                doc.contenu_brouillon = request.POST.get('contenu', '').strip()
                doc.statut = 'valide'
                doc.date_validation = timezone.now()
                doc.valide_par = request.user
                doc.save()
                JournalDocumentGenere.objects.create(document=doc, utilisateur=request.user, action='validation')
                messages.success(request, "Document validé.")
            else:
                messages.error(request, "Un document envoyé ne peut plus être modifié.")
        elif action == 'rejeter':
            doc.statut = 'rejete'
            doc.save()
            JournalDocumentGenere.objects.create(document=doc, utilisateur=request.user, action='rejet')
            try:
                parametres = json.loads(doc.parametres_generation or '{}')
                contenu = generer_mise_en_demure(parametres.get('gabarit_rempli', doc.contenu_brouillon),
                                                 parametres.get('valeurs', {}))
                doc.contenu_brouillon = contenu
                doc.statut = 'brouillon'
                doc.save()
                JournalDocumentGenere.objects.create(document=doc, utilisateur=request.user, action='generation')
                messages.success(request, "Brouillon rejeté et régénéré.")
            except Exception as e:
                messages.error(request, "Rejet enregistré. Régénération impossible : {}".format(_message_ia_securise(e)))
        return redirect('reviser_document', pk=doc.pk)
    return render(request, 'documents_revision.html', {'doc': doc, 'active': 'documents_generes'})


@login_required
def liste_documents_generes(request):
    qs = (DocumentGenere.objects.select_related('dossier', 'valide_par')
          .filter(dossier__avocat=request.user).order_by('-date_generation'))
    paginator = Paginator(qs, 10)
    documents = paginator.get_page(request.GET.get('page', 1))
    return render(request, 'documents_generes_liste.html', {'documents': documents, 'active': 'documents_generes'})


def _document_accessible(doc, user):
    return (doc.auteur == user) or (doc.dossier is not None and doc.dossier.avocat == user)


def _contexte_document(doc):
    dossier = doc.dossier
    client = dossier.client if dossier else doc.client
    return (
        f"Client : {client if client else 'non renseigné'}\n"
        f"Dossier : {dossier.titre if dossier else 'non renseigné'}\n"
        f"Type de procédure : {dossier.get_type_procedure_display() if dossier else '—'}\n"
        f"Notes du dossier : {dossier.notes if dossier else '—'}\n"
        f"Document analysé : {doc.nom_original} ({doc.type_fichier})"
    )


@login_required
def analyser_document(request, document_id):
    doc = get_object_or_404(Document.objects.select_related('dossier', 'client'), pk=document_id)
    if not _document_accessible(doc, request.user):
        messages.error(request, "Vous n'avez pas accès à ce document.")
        return redirect('liste_documents')
    if doc.type_fichier not in ('docx', 'pdf', 'txt'):
        messages.error(request, "Ce type de fichier ne peut pas être analysé.")
        return redirect('liste_documents')
    if quota_restant(request.user) <= 0:
        messages.error(request, "Quota IA mensuel dépassé. Contactez l'administrateur.")
        return redirect('liste_documents')
    try:
        texte = extraire_texte(doc.fichier.path, doc.fichier.name)
    except Exception as e:
        messages.error(request, "Impossible de lire le fichier : {}".format(_message_erreur_securise(e)))
        return redirect('liste_documents')
    if not texte.strip():
        messages.error(request, "Aucun texte extractible (PDF scanné ?).")
        return redirect('liste_documents')
    try:
        resultat = analyser_document_ia(texte, _contexte_document(doc))
    except Exception as e:
        messages.error(request, "Erreur lors de l'analyse IA : {}".format(_message_ia_securise(e)))
        return redirect('liste_documents')
    analyse = AnalyseDocument.objects.create(
        document=doc, contenu_analyse=resultat, auteur=request.user)
    enregistrer_usage(request.user, 'analyse', octets=len(texte))
    messages.success(request, "Analyse générée avec succès.")
    return redirect('fiche_analyse', analyse_id=analyse.pk)


@login_required
def fiche_analyse(request, analyse_id):
    analyse = get_object_or_404(
        AnalyseDocument.objects.select_related('document', 'document__dossier', 'auteur'),
        pk=analyse_id)
    if not _document_accessible(analyse.document, request.user):
        messages.error(request, "Vous n'avez pas accès à cette analyse.")
        return redirect('liste_documents')
    return render(request, 'documents_analyse.html', {'analyse': analyse, 'active': 'documents'})


def _nom_fichier_export(doc):
    client = doc.dossier.client
    nom_client = f"{client.nom}_{client.prenom}".replace(' ', '_').lower()
    date = doc.date_validation or doc.date_generation
    return f"mise_en_demure_{nom_client}_{date.strftime('%Y-%m-%d')}"


def _exporter_docx(doc):
    document = DocxDocument()
    for para in doc.contenu_brouillon.splitlines():
        if para.strip():
            document.add_paragraph(para)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = 'attachment; filename="{}.docx"'.format(_nom_fichier_export(doc))
    document.save(response)
    return response


def _exporter_pdf(doc):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="{}.pdf"'.format(_nom_fichier_export(doc))
    document = SimpleDocTemplate(response, pagesize=letter)
    styles = getSampleStyleSheet()
    style_normal = styles['Normal']
    style_normal.fontName = 'Helvetica'
    style_normal.fontSize = 11
    style_normal.leading = 15
    elements = []
    for para in doc.contenu_brouillon.splitlines():
        ligne = para.strip()
        if ligne:
            texte = ligne.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            elements.append(Paragraph(texte, style_normal))
            elements.append(Spacer(1, 6))
    document.build(elements)
    return response


@login_required
def exporter_document(request, pk, format):
    doc = get_object_or_404(DocumentGenere, pk=pk, dossier__avocat=request.user)
    if doc.statut != 'valide':
        messages.error(request, "Export autorisé uniquement pour un document validé.")
        return redirect('reviser_document', pk=doc.pk)
    if format == 'docx':
        return _exporter_docx(doc)
    if format == 'pdf':
        return _exporter_pdf(doc)
    return redirect('reviser_document', pk=doc.pk)


def _lignes_analyse_brutes(texte):
    """Convertit le markdown de l'analyse en texte brut (paragraphes)."""
    lignes = []
    for ligne in (texte or '').splitlines():
        l = ligne.strip()
        if not l:
            lignes.append('')
            continue
        l = re.sub(r'^#{1,6}\s+', '', l)
        l = re.sub(r'^\s*[-*]\s+', '- ', l)
        l = re.sub(r'^\s*\d+[.)]\s+', '', l)
        l = l.replace('**', '').replace('__', '').replace('*', '').replace('`', '')
        lignes.append(l)
    return lignes


def _nom_fichier_analyse(analyse):
    base = re.sub(r'[^\w\-.]+', '_', (analyse.document.nom_original or '').rsplit('.', 1)[0])
    return f"analyse_{base}_{analyse.date_analyse.strftime('%Y-%m-%d')}"


def _exporter_analyse_docx(analyse):
    document = DocxDocument()
    document.add_heading(analyse.document.nom_original, level=1)
    document.add_paragraph("Analyse du {} — {}".format(
        analyse.date_analyse.strftime('%d/%m/%Y %H:%M'),
        analyse.auteur.get_full_name() or analyse.auteur.username if analyse.auteur else '—'))
    for para in _lignes_analyse_brutes(analyse.contenu_analyse):
        document.add_paragraph(para)
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = 'attachment; filename="{}.docx"'.format(_nom_fichier_analyse(analyse))
    document.save(response)
    return response


def _exporter_analyse_pdf(analyse):
    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="{}.pdf"'.format(_nom_fichier_analyse(analyse))
    document = SimpleDocTemplate(response, pagesize=letter)
    styles = getSampleStyleSheet()
    style_normal = styles['Normal']
    style_normal.fontName = 'Helvetica'
    style_normal.fontSize = 11
    style_normal.leading = 15
    elements = []
    elements.append(Paragraph(analyse.document.nom_original.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'), styles['Title']))
    for para in _lignes_analyse_brutes(analyse.contenu_analyse):
        ligne = para.strip()
        if ligne:
            texte = ligne.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            elements.append(Paragraph(texte, style_normal))
            elements.append(Spacer(1, 6))
    document.build(elements)
    return response


@login_required
def exporter_analyse(request, analyse_id, format):
    analyse = get_object_or_404(
        AnalyseDocument.objects.select_related('document', 'document__dossier', 'auteur'),
        pk=analyse_id)
    if not _document_accessible(analyse.document, request.user):
        messages.error(request, "Vous n'avez pas accès à cette analyse.")
        return redirect('liste_documents')
    if format == 'docx':
        return _exporter_analyse_docx(analyse)
    if format == 'pdf':
        return _exporter_analyse_pdf(analyse)
    return redirect('fiche_analyse', analyse_id=analyse.pk)